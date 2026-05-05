from docx import Document
from data import case_data

doc = Document()
doc.add_heading("LEGAL NOTICE", level=1)

doc.add_paragraph(
    f"This legal notice is issued on behalf of {case_data['petitioner']} "
    f"regarding an accident dated {case_data['accident_date']}."
)

doc.add_paragraph(
    "You are hereby called upon to settle the claim failing which "
    "legal proceedings shall be initiated."
)

doc.save("output/word/legal_notice.docx")
print("Word file generated")
